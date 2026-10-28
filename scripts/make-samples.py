"""
Builds the sample study materials in samples/.

All sample content was written for this project (it is NOT real course material),
so it is safe to keep in a public repository. The "course policy" facts are invented on
purpose: general knowledge can't answer them, so they prove answers come from the documents.

Requires: pip install reportlab python-docx python-pptx fonttools
Run from the project root:  python scripts/make-samples.py
"""
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak

OUT = Path(__file__).resolve().parent.parent / "samples"
OUT.mkdir(exist_ok=True)
NOTICE = "SAMPLE CONTENT: written for the Gen AI demo. Not real course material."


def english_pdf():
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=11, leading=15)
    h1, h2 = styles["Title"], styles["Heading2"]
    pages = [
        [
            Paragraph("FIN101 Lecture 3: The Time Value of Money", h1),
            Paragraph(NOTICE, body),
            Spacer(1, 10),
            Paragraph("1. Why money has a time value", h2),
            Paragraph(
                "A baht today is worth more than a baht received in the future. Money available now can be invested to earn "
                "interest, and future payments carry risk and lose purchasing power to inflation. The time value of money (TVM) "
                "is the idea that we must adjust for timing before comparing cash flows that occur on different dates.",
                body,
            ),
            Paragraph(
                "Two operations do this adjustment. Compounding moves a value forward in time to find its future value. "
                "Discounting moves a value back in time to find its present value. The rate used is called the interest rate "
                "when compounding and the discount rate when discounting.",
                body,
            ),
        ],
        [
            Paragraph("2. Future value and compound interest", h2),
            Paragraph(
                "With simple interest, interest is earned only on the original principal. With compound interest, interest is "
                "also earned on previously earned interest. The future value of a single amount is FV = PV x (1 + r)^n, where "
                "PV is the present value, r is the interest rate per period and n is the number of periods.",
                body,
            ),
            Paragraph(
                "Worked example: 10,000 baht invested at 5% per year for 3 years grows to 10,000 x 1.05^3 = 11,576.25 baht. "
                "With simple interest it would grow to only 11,500 baht. The extra 76.25 baht is interest earned on interest.",
                body,
            ),
            Paragraph(
                "When interest is compounded m times per year, use the periodic rate r/m and the number of periods n x m. "
                "More frequent compounding gives a higher future value. The effective annual rate is EAR = (1 + r/m)^m - 1.",
                body,
            ),
        ],
        [
            Paragraph("3. Present value and discounting", h2),
            Paragraph(
                "The present value of a future amount is PV = FV / (1 + r)^n. A higher discount rate or a longer waiting time "
                "makes the present value smaller. For example, 20,000 baht to be received in 4 years, discounted at 6% per "
                "year, is worth 20,000 / 1.06^4 = 15,841.87 baht today.",
                body,
            ),
            Paragraph(
                "Common mistake: using the annual rate with monthly periods. If cash flows are monthly, convert both the rate "
                "and the number of periods to months before applying the formula.",
                body,
            ),
            Paragraph(
                "Annuities are equal payments made every period. The present value of an ordinary annuity is "
                "PV = PMT x [1 - (1 + r)^-n] / r. A loan repayment schedule is an annuity viewed from the lender's side.",
                body,
            ),
        ],
        [
            Paragraph("4. FIN101 course policies (invented for the demo)", h2),
            Paragraph(
                "Problem sets are due every Friday at 17:00. A late problem set loses 15% of its marks for each day it is late, "
                "and problem sets more than 3 days late receive zero marks.",
                body,
            ),
            Paragraph(
                "The midterm exam covers Lectures 1 to 4 and allows one A4 formula sheet written by hand on both sides. "
                "Financial calculators are allowed; phones are not.",
                body,
            ),
            Paragraph(
                "Office hours are on Wednesdays from 13:00 to 15:00 in room FIN-204, or online by appointment.",
                body,
            ),
        ],
    ]
    story = []
    for i, page in enumerate(pages):
        story.extend(page)
        if i < len(pages) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(str(OUT / "fin101-time-value-of-money.pdf"), pagesize=A4, title="FIN101 Lecture 3: The Time Value of Money").build(story)


def english_pptx():
    prs = Presentation()
    slides = [
        ("ECON102: Supply and Demand", NOTICE, "This deck introduces how prices are set in a competitive market."),
        (
            "The law of demand",
            "When the price of a good rises, the quantity demanded falls, other things equal.\nThe demand curve slopes downward.",
            "Two reasons: the substitution effect (buyers switch to cheaper goods) and the income effect (higher prices reduce real income).",
        ),
        (
            "The law of supply",
            "When the price of a good rises, the quantity supplied rises, other things equal.\nThe supply curve slopes upward.",
            "Higher prices make production more profitable, so firms produce more and new firms may enter.",
        ),
        (
            "Market equilibrium",
            "Equilibrium is where quantity demanded equals quantity supplied.\nAbove equilibrium there is a surplus; below it there is a shortage.",
            "Surpluses push prices down and shortages push prices up until the market clears.",
        ),
        (
            "Shifts vs. movements",
            "A change in the good's own price causes a movement along the curve.\nA change in income, tastes, input costs or related prices shifts the curve.",
            "Common exam mistake: saying 'demand increased' when only the quantity demanded changed because the price fell.",
        ),
        (
            "ECON102 tutorial rule (invented for the demo)",
            "Tutorial worksheets must be uploaded by Sunday 20:00.\nThe best 8 of 10 worksheets count towards the final grade.",
            "",
        ),
    ]
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        tf = slide.placeholders[1].text_frame
        lines = body.split("\n")
        tf.text = lines[0]
        for line in lines[1:]:
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(20)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    prs.save(OUT / "econ102-supply-and-demand.pptx")


def thai_docx():
    doc = Document()
    doc.add_heading("FIN101 บทที่ 4: ดอกเบี้ยทบต้นและการวางแผนการออม", level=1)
    doc.add_paragraph("เนื้อหาตัวอย่าง: เขียนขึ้นสำหรับการสาธิต Gen AI ไม่ใช่เอกสารประกอบการเรียนจริง")
    doc.add_heading("1. ดอกเบี้ยทบต้นคืออะไร", level=2)
    doc.add_paragraph(
        "ดอกเบี้ยทบต้น (compound interest) คือดอกเบี้ยที่คิดจากเงินต้นรวมกับดอกเบี้ยที่ได้รับในงวดก่อนหน้า "
        "ต่างจากดอกเบี้ยเชิงเดียว (simple interest) ที่คิดจากเงินต้นเพียงอย่างเดียว "
        "ยิ่งระยะเวลานานขึ้น ผลของดอกเบี้ยทบต้นก็ยิ่งมากขึ้น"
    )
    doc.add_paragraph("สูตรมูลค่าในอนาคต: FV = PV × (1 + r)^n โดย PV คือเงินต้น r คืออัตราดอกเบี้ยต่องวด และ n คือจำนวนงวด")
    doc.add_heading("2. ตัวอย่างการคำนวณ", level=2)
    doc.add_paragraph(
        "ถ้าฝากเงิน 50,000 บาท อัตราดอกเบี้ย 4% ต่อปี เป็นเวลา 5 ปี แบบทบต้นปีละครั้ง "
        "จะได้เงินรวมประมาณ 50,000 × 1.04^5 = 60,832.65 บาท "
        "ส่วนดอกเบี้ยเชิงเดียวจะได้เพียง 60,000 บาท"
    )
    doc.add_heading("3. กฎ 72", level=2)
    doc.add_paragraph(
        "กฎ 72 (Rule of 72) ใช้ประมาณจำนวนปีที่เงินจะเพิ่มเป็นสองเท่า โดยนำ 72 หารด้วยอัตราดอกเบี้ยต่อปี (เป็นเปอร์เซ็นต์) "
        "เช่น อัตรา 6% ต่อปี เงินจะเพิ่มเป็นสองเท่าในประมาณ 72 / 6 = 12 ปี "
        "ข้อผิดพลาดที่พบบ่อยคือการใช้อัตราดอกเบี้ยในรูปทศนิยม เช่น 0.06 แทน 6"
    )
    doc.add_heading("4. ข้อกำหนดของรายวิชา (สมมติขึ้นเพื่อการสาธิต)", level=2)
    doc.add_paragraph(
        "การสอบย่อยครั้งที่ 2 ของวิชา FIN101 จัดในสัปดาห์ที่ 9 และออกสอบเฉพาะบทที่ 3 และบทที่ 4 "
        "นักศึกษาสามารถนำเครื่องคิดเลขทางการเงินเข้าห้องสอบได้"
    )
    doc.save(OUT / "fin101-compound-interest-thai.docx")


def markdown_notes():
    (OUT / "study-skills-notes.md").write_text(
        f"""# Study skills: active recall and spaced repetition

{NOTICE}

## Active recall
Active recall means testing yourself instead of re-reading. Closing the book and writing down everything you remember
forces your brain to retrieve the information, which strengthens memory more than passive review.

## Spaced repetition
Spaced repetition means reviewing material at increasing intervals, for example after 1 day, 3 days, 7 days and 14 days.
Reviewing just before you would forget is more efficient than cramming everything the night before an exam.

## Interleaving
Interleaving means mixing different types of problems in one study session instead of practising one type at a time.
It feels harder, but it improves your ability to choose the right method in an exam.

## A simple weekly routine
1. After each lecture, write a one-page summary from memory.
2. Two days later, answer five practice questions without notes.
3. At the end of the week, review only the questions you got wrong.
""",
        encoding="utf-8",
    )


if __name__ == "__main__":
    english_pdf()
    english_pptx()
    thai_docx()
    markdown_notes()
    for p in sorted(OUT.iterdir()):
        print(f"{p.name}  {p.stat().st_size:,} bytes")
